import json
import os
import sys
from decimal import Decimal

import docx
import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "samples"))

import make_samples  # noqa: E402
from fsreview import checks, pipeline  # noqa: E402
from fsreview.classify import analyse_table, classify_document, classify_title  # noqa: E402
from fsreview.extract import read_docx  # noqa: E402
from fsreview.model import Document, Row, Table  # noqa: E402
from fsreview.numbers import format_amount, parse_amount  # noqa: E402
from fsreview.pdf_engine import find_office  # noqa: E402

FAST = {"financial_year": "2025", "entity_name": "Northwind Trading Limited", "category": "MSC",
        "outputs": ["docx", "xlsx", "report"]}


@pytest.fixture(scope="module")
def samples(tmp_path_factory):
    d = tmp_path_factory.mktemp("samples")
    make_samples.make_docx(str(d / "fs.docx"))
    make_samples.make_xlsx(str(d / "fs.xlsx"))
    make_samples.make_txt(str(d / "fs.txt"))
    return d


@pytest.fixture(scope="module")
def docx_run(samples, tmp_path_factory):
    out = tmp_path_factory.mktemp("out")
    data = pipeline.run(str(samples / "fs.docx"), str(out), dict(FAST), progress=lambda m: None)
    return data, out


def messages(data, category=None, status=None):
    return [i["message"] for i in data["issues"]
            if (category is None or i["category"] == category) and (status is None or i["status"] == status)]


# ------------------------------------------------------------------ numbers

@pytest.mark.parametrize("text,value", [
    ("1,234", Decimal(1234)), ("(1,234)", Decimal(-1234)), ("-1,234", Decimal(-1234)), ("1234.50", Decimal("1234.50")),
    ("-", Decimal(0)), ("nil", Decimal(0)), ("–", Decimal(0)), ("$ 500", Decimal(500)), ("(500", None),
    ("12.5%", None), ("Revenue", None), ("", None),
])
def test_parse_amount(text, value):
    assert parse_amount(text) == value


def test_format_amount():
    assert format_amount(Decimal(-1234)) == "(1,234)"
    assert format_amount(Decimal(-1234), negative="minus") == "-1,234"
    assert format_amount(Decimal(0)) == "-"
    assert format_amount(Decimal("1234.5"), 2) == "1,234.50"


def test_titles_and_contents_entries():
    assert classify_title("STATEMENT OF FINANCIAL POSITION") == "SOFP"
    assert classify_title("Statement of Profit or Loss and Other Comprehensive Income") == "SOCI"
    assert classify_title("Statement of cash flows\t6") == "CONTENTS_ENTRY"
    assert classify_title("Statement of financial position .......... 3") == "CONTENTS_ENTRY"
    assert classify_title("Revenue increased in note 12") is None


# ------------------------------------------------------------------ casting

def _table(rows, section="SOFP"):
    t = Table(rows=[Row(cells=list(r)) for r in rows], section=section)
    return analyse_table(t)


def test_casting_finds_wrong_total():
    t = _table([("", "2025", "2024"), ("Inventories", "100", "90"), ("Cash", "50", "40"), ("Total current assets", "160", "130")])
    doc = Document(blocks=[t])
    found = checks.cast_tables(doc)
    errs = [i for i in found if i.severity == "error"]
    assert len(errs) == 1 and "150" in errs[0].message and "2025" in errs[0].message


def test_casting_accepts_layered_totals():
    t = _table([("", "2025"), ("PPE", "100"), ("", "100"), ("Inventory", "30"), ("Cash", "20"), ("", "50"),
                ("Total assets", "150"), ("Share capital", "100"), ("Total equity", "100"), ("Borrowings", "20"),
                ("Current liabilities", ""), ("Payables", "30"), ("", "30"), ("Total liabilities", "50"),
                ("Total equity and liabilities", "150")])
    found = checks.cast_tables(Document(blocks=[t]))
    assert not [i for i in found if i.severity == "error"], [i.message for i in found]


def test_running_totals_in_profit_or_loss():
    t = _table([("", "2025"), ("Revenue", "1,000"), ("Cost of sales", "(600)"), ("Gross profit", "400"),
                ("Expenses", "(150)"), ("Operating profit", "250"), ("Tax", "(50)"), ("Profit for the year", "200")],
               section="SOCI")
    found = checks.cast_tables(Document(blocks=[t]))
    assert not [i for i in found if i.severity == "error"]


# ------------------------------------------------------------- full review

def test_word_review_finds_planted_problems(docx_run):
    data, _ = docx_run
    msgs = " | ".join(messages(data))
    assert "refers to note 11, but there is no note 11" in msgs
    assert "Invalid date: '31 June 2025'" in msgs
    assert "IAS 18" in msgs
    assert "'recievables' → 'receivables'" in msgs
    assert "Repeated word: 'the the'" in msgs
    assert "Different period-end dates" in msgs
    assert "balance sheet" in msgs.lower()
    assert not messages(data, "casting") or all("cast" in m for m in messages(data, "casting"))


def test_word_review_cross_checks_pass(docx_run):
    data, _ = docx_run
    passed = " | ".join(messages(data, status="passed"))
    assert "total assets vs total equity and liabilities: agrees at 22,000" in passed
    assert "Closing total equity in the statement of changes in equity vs the statement of financial position" in passed
    assert "Profit for the year in the statement of changes in equity vs the statement of profit or loss" in passed
    assert "Opening cash + net change = closing cash (current year): agrees" in passed


def test_sections_identified(docx_run):
    data, _ = docx_run
    for s in ("Statement of Financial Position", "Statement of Profit or Loss and Other Comprehensive Income",
              "Statement of Changes in Equity", "Statement of Cash Flows", "Notes to the Financial Statements",
              "Significant Accounting Policies", "Directors' / Management Report"):
        assert s in data["sections"]


def test_formatting_preserves_content(docx_run, samples):
    data, out = docx_run
    integrity = [i for i in data["issues"] if i["category"] == "integrity"]
    assert integrity and all(i["status"] == "passed" for i in integrity), integrity
    before = read_docx(str(samples / "fs.docx"))
    after = read_docx(str(out / "fs - formatted.docx"))
    vals = lambda m: [parse_amount(c) for t in m.tables() for r in t.rows for c in r.cells if parse_amount(c) is not None]
    assert vals(before) == vals(after)


def test_figures_presented_consistently(docx_run):
    data, out = docx_run
    edits = {(e["before"], e["after"]) for e in data["changes"]["edits"]}
    assert ("12500", "12,500") in edits
    assert ("-5,200", "(5,200)") in edits
    d = docx.Document(str(out / "fs - formatted.docx"))
    sec = d.sections[0]
    assert round(sec.page_width.cm, 1) == 21.0
    assert any("PAGE" in (el.text or "") for el in sec.footer._element.iter(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instrText"))


def test_spelling_fix_option(samples, tmp_path):
    s = dict(FAST, apply_spelling=True, outputs=["docx", "report"])
    data = pipeline.run(str(samples / "fs.docx"), str(tmp_path), s, progress=lambda m: None)
    text = "\n".join(p.text for p in docx.Document(str(tmp_path / "fs - formatted.docx")).paragraphs)
    assert "depreciaton" not in text and "accumulated depreciation" in text
    assert all(i["status"] == "passed" for i in data["issues"] if i["category"] == "integrity")
    assert any(e["reason"] == "Spelling corrected" for e in data["changes"]["edits"])


def test_casting_error_in_word(samples, tmp_path):
    src = tmp_path / "bad.docx"
    d = docx.Document(str(samples / "fs.docx"))
    for t in d.tables:
        for row in t.rows:
            if row.cells[0].text.strip() == "Total assets":
                row.cells[2].paragraphs[0].runs[0].text = "22,100"
    d.save(str(src))
    data = pipeline.run(str(src), str(tmp_path / "o"), dict(FAST, outputs=["report"]), progress=lambda m: None)
    errs = [i["message"] for i in data["issues"] if i["severity"] == "error"]
    assert any("Total assets" in m and "22,100" in m for m in errs)
    assert any("total assets vs total equity and liabilities" in m for m in errs)


def test_excel_workbook_untouched_except_print_setup(samples, tmp_path):
    import zipfile
    data = pipeline.run(str(samples / "fs.xlsx"), str(tmp_path), dict(FAST), progress=lambda m: None)
    wb = openpyxl.load_workbook(str(tmp_path / "fs - formatted.xlsx"))
    assert wb["SOFP"]["C10"].value == "=SUM(C8:C9)"
    assert wb["SOFP"]["C16"].value == "=C10+C15"
    assert wb["SOFP"].page_setup.fitToWidth == 1 and wb["SOFP"].print_options.horizontalCentered
    assert any("byte-for-byte unchanged" in m for m in messages(data, "integrity", "passed"))
    with zipfile.ZipFile(samples / "fs.xlsx") as a, zipfile.ZipFile(tmp_path / "fs - formatted.xlsx") as b:
        assert a.read("xl/styles.xml") == b.read("xl/styles.xml")
    assert (tmp_path / "fs - financial analysis.xlsx").exists()


def test_audit_pack_workbook(tmp_path):
    import zipfile
    from fsreview.extract import read_document
    src = tmp_path / "pack.xlsx"
    make_samples.make_audit_pack(str(src))
    m = classify_document(read_document(str(src)))
    assert m.sheets["statements"] == ["BS", "Notes"]
    assert m.sheets["working"] == ["TB 2026"] and m.sheets["hidden"] == ["BS (old)"]
    assert any("decimals that the number format hides" in w for w in m.warnings)
    text = m.text()
    assert "check" not in text                         # outside the print area
    heads = [p.text for p in m.paragraphs() if p.role == "note_heading"]
    assert "4.1 Allocation of depreciation" in heads
    data = pipeline.run(str(src), str(tmp_path / "o"), dict(FAST, financial_year="2026"), progress=lambda m: None)
    msgs = " | ".join(messages(data))
    assert "'SAHRE' → 'SHARE'" in msgs
    assert "no third column as at the beginning of the preceding period" in msgs
    assert "the last note in the financial statements is note 4" in msgs
    assert any("agrees at 1,750,001" in x for x in messages(data, "accounting", "passed"))
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(tmp_path / "o" / "pack - formatted.xlsx") as b:
        media = [n for n in a.namelist() if n.startswith("xl/media/")]
        assert media and all(a.read(n) == b.read(n) for n in media)       # the logo survives
        assert a.read("xl/worksheets/sheet3.xml") == b.read("xl/worksheets/sheet3.xml")   # working paper


def test_text_input(samples, tmp_path):
    data = pipeline.run(str(samples / "fs.txt"), str(tmp_path), dict(FAST), progress=lambda m: None)
    assert "Statement of Financial Position" in data["sections"]
    assert any("agrees at 22,000" in m for m in messages(data, "accounting", "passed"))
    assert all(i["status"] == "passed" for i in data["issues"] if i["category"] == "integrity")


@pytest.mark.skipif(not find_office(), reason="LibreOffice is needed to make a PDF test input")
def test_pdf_input(samples, tmp_path):
    import subprocess
    subprocess.run([find_office(), "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path),
                    str(samples / "fs.docx")], check=True, capture_output=True, timeout=180)
    data = pipeline.run(str(tmp_path / "fs.pdf"), str(tmp_path / "o"), dict(FAST), progress=lambda m: None)
    msgs = " | ".join(messages(data))
    assert "no note 11" in msgs and "31 June 2025" in msgs
    assert not [i for i in data["issues"] if i["category"] == "casting" and i["severity"] == "error"]
    assert any("agrees at 14,600" in m for m in messages(data, "accounting", "passed"))


def test_report_files(docx_run):
    data, out = docx_run
    assert (out / "fs - review report.docx").exists()
    j = json.loads((out / "fs - review report.json").read_text(encoding="utf-8"))
    assert j["summary"]["error"] == data["summary"]["error"] > 0


# ------------------------------------------------- entity category and disclosures

def _model(path):
    from fsreview.extract import read_document
    return classify_document(read_document(str(path)))


def test_category_suggestions(samples, tmp_path):
    from fsreview import entity
    from fsreview.figures import extract_figures
    m = _model(samples / "fs.docx")
    assert entity.suggest(m, extract_figures(m))["suggested"] == "MSC"
    m.blocks[0].text = "NORTHWIND TRADING (PRIVATE) LIMITED"
    assert entity.suggest(m, extract_figures(m))["suggested"] == "SSC"
    m.blocks[0].text = "The shares of the company are listed on the Pakistan Stock Exchange."
    assert entity.suggest(m, extract_figures(m))["suggested"] == "PIC"
    figs = extract_figures(m)
    figs["revenue"] = (Decimal(2_000_000), Decimal(1_500_000))      # Rs '000 -> 1.5 billion
    m.blocks[0].text = "NORTHWIND"
    assert entity.suggest(m, figs)["suggested"] == "LSC"


def test_disclosure_checklist(samples):
    from fsreview import disclosures
    m = _model(samples / "fs.docx")
    rows, issues = disclosures.run_checklist(m, "MSC")
    st = {r["id"]: r["status"] for r in rows}
    assert st["ppe"] == "present" and st["share_capital"] == "present"
    assert st["related_parties"] == "missing" and st["authorisation"] == "missing"
    assert st["deferred_tax"] == "n/a"
    assert "eps" not in st                       # not required below PIC
    rows, _ = disclosures.run_checklist(m, "PIC", {"listed": True})
    assert {r["id"]: r["status"] for r in rows}["eps"] == "missing"
    assert any("Disclosure not found: Transactions and balances with related parties" in i.message for i in issues)


def test_reference_library(samples, tmp_path, monkeypatch):
    from fsreview import library
    monkeypatch.setattr(library, "LIB", str(tmp_path / "lib"))
    ref = tmp_path / "icap.docx"
    d = docx.Document()
    d.add_paragraph("Notes to the Financial Statements")
    for i, t in enumerate(["The company and its operations", "Trade debts", "Long term financing",
                           "Number of employees"], 1):
        d.add_paragraph(f"{i}. {t}").runs[0].bold = True
        d.add_paragraph("Text.")
    d.save(str(ref))
    meta = library.add(str(ref), "ICAP MSC", "reference", "MSC")
    assert len(meta["note_headings"]) == 4
    rows, issues = library.compare(_model(samples / "fs.docx"), "MSC")
    st = {r["note"]: r["status"] for r in rows}
    assert st["The company and its operations"] == "present"      # matched to 'General information'
    assert st["Trade debts"] == "present"                         # matched to 'Trade and other receivables'
    assert st["Long term financing"] == "missing" and st["Number of employees"] == "missing"
    assert library.compare(_model(samples / "fs.docx"), "PIC") == ([], [])


# ------------------------------------------------------------- analysis and layout

def test_ratio_analysis(samples):
    from fsreview import analysis
    r = analysis.compute(_model(samples / "fs.docx"))
    by = {x["name"]: x for x in r["ratios"]}
    assert round(by["Gross profit margin"]["cy"], 1) == 37.0        # 10,500 / 28,400
    assert round(by["Current ratio"]["cy"], 2) == round(8250 / 3900, 2)
    assert round(by["Debt to equity"]["cy"], 2) == round(3500 / 14600, 2)
    assert any(g["name"] == "Revenue" and round(g["pct"], 1) == 9.7 for g in r["growth"])


def test_full_outputs_with_analysis_and_big4_layout(docx_run):
    data, out = docx_run
    assert (out / "fs - financial analysis.docx").exists()
    assert data["charts"] and all((out / c["file"]).exists() for c in data["charts"])
    wb = openpyxl.load_workbook(str(out / "fs - formatted.xlsx"))
    assert "Ratios" in wb.sheetnames and "Analysis" in wb.sheetnames
    d = docx.Document(str(out / "fs - formatted.docx"))
    cells = " ".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "Chief Executive" in cells and "Director" in cells
    assert "------------ $'000 ------------" in cells
    text = "\n".join(p.text for p in d.paragraphs)
    assert "5\tPROPERTY, PLANT AND EQUIPMENT" in text or "5.\tPROPERTY, PLANT AND EQUIPMENT" in text
    assert data["category"]["code"] == "MSC" and data["checklist"]
