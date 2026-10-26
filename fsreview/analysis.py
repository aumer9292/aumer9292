"""Analytical review: ratios, horizontal and vertical analysis, and charts.

Everything is computed from the statements as presented (current year against
the comparative). Each ratio carries its formula and a one-line reading so a
reviewer can see what moved and by how much; the commentary is mechanical and
is a starting point for the reviewer's own analytical procedures.
"""

import os
from decimal import Decimal

from .classify import SECTIONS, is_financial_table, row_values
from .figures import extract_figures

CY_COLOR, PY_COLOR = "#2a78d6", "#eb6834"      # validated categorical slots 1 and 2
INK, MUTED, GRID = "#0b0b0b", "#52514e", "#e4e3df"


def _f(figs, key, i):
    v = figs.get(key)
    return float(v[i]) if v and v[i] is not None else None


def _div(a, b):
    if a is None or b in (None, 0):
        return None
    return a / b


RATIOS = [
    # group, name, unit, formula text, function(figs, i), better ("up"/"down"/None)
    ("Profitability", "Gross profit margin", "%", "Gross profit ÷ revenue",
     lambda f, i: _pct(_div(_f(f, "gross_profit", i), _f(f, "revenue", i))), "up"),
    ("Profitability", "Operating profit margin", "%", "Operating profit ÷ revenue",
     lambda f, i: _pct(_div(_f(f, "operating_profit", i), _f(f, "revenue", i))), "up"),
    ("Profitability", "Net profit margin", "%", "Profit after tax ÷ revenue",
     lambda f, i: _pct(_div(_f(f, "profit", i), _f(f, "revenue", i))), "up"),
    ("Profitability", "Return on equity", "%", "Profit after tax ÷ total equity",
     lambda f, i: _pct(_div(_f(f, "profit", i), _f(f, "total_equity", i))), "up"),
    ("Profitability", "Return on assets", "%", "Profit after tax ÷ total assets",
     lambda f, i: _pct(_div(_f(f, "profit", i), _f(f, "total_assets", i))), "up"),
    ("Profitability", "Effective tax rate", "%", "Tax expense ÷ profit before tax",
     lambda f, i: _pct(_div(_abs(_f(f, "tax", i)), _f(f, "pbt", i))), None),
    ("Liquidity", "Current ratio", "times", "Current assets ÷ current liabilities",
     lambda f, i: _div(_f(f, "total_current_assets", i), _f(f, "total_current_liabilities", i)), "up"),
    ("Liquidity", "Quick ratio", "times", "(Current assets − inventories − stores) ÷ current liabilities",
     lambda f, i: _div(_sub(_f(f, "total_current_assets", i), _f(f, "inventory", i), _f(f, "stores", i)),
                       _f(f, "total_current_liabilities", i)), "up"),
    ("Liquidity", "Cash ratio", "times", "Cash and bank ÷ current liabilities",
     lambda f, i: _div(_f(f, "cash", i), _f(f, "total_current_liabilities", i)), "up"),
    ("Efficiency", "Inventory days", "days", "Inventories ÷ cost of sales × 365",
     lambda f, i: _days(_f(f, "inventory", i), _abs(_f(f, "cost_of_sales", i))), "down"),
    ("Efficiency", "Receivable days", "days", "Trade debts ÷ revenue × 365",
     lambda f, i: _days(_f(f, "receivables", i), _f(f, "revenue", i)), "down"),
    ("Efficiency", "Payable days", "days", "Trade and other payables ÷ cost of sales × 365",
     lambda f, i: _days(_f(f, "payables", i), _abs(_f(f, "cost_of_sales", i))), None),
    ("Efficiency", "Asset turnover", "times", "Revenue ÷ total assets",
     lambda f, i: _div(_f(f, "revenue", i), _f(f, "total_assets", i)), "up"),
    ("Leverage", "Debt to equity", "times", "Borrowings and lease liabilities ÷ total equity",
     lambda f, i: _div(_f(f, "total_debt", i), _f(f, "total_equity", i)), "down"),
    ("Leverage", "Gearing", "%", "Debt ÷ (debt + equity)",
     lambda f, i: _pct(_div(_f(f, "total_debt", i), _add(_f(f, "total_debt", i), _f(f, "total_equity", i)))), "down"),
    ("Leverage", "Liabilities to assets", "%", "Total liabilities ÷ total assets",
     lambda f, i: _pct(_div(_f(f, "total_liabilities", i), _f(f, "total_assets", i))), "down"),
    ("Leverage", "Interest cover", "times", "Operating profit ÷ finance cost",
     lambda f, i: _div(_f(f, "operating_profit", i) or _add(_f(f, "pbt", i), _abs(_f(f, "finance_cost", i))),
                       _abs(_f(f, "finance_cost", i))), "up"),
    ("Cash flow", "Cash flow to profit", "times", "Net cash from operating activities ÷ profit after tax",
     lambda f, i: _div(_f(f, "cfo", i), _f(f, "profit", i)), "up"),
    ("Cash flow", "Free cash flow", "amount", "Operating cash flow − capital expenditure",
     lambda f, i: _sub(_f(f, "cfo", i), _abs(_f(f, "capex", i))) if _f(f, "cfo", i) is not None else None, "up"),
]


def _pct(v):
    return v * 100 if v is not None else None


def _abs(v):
    return abs(v) if v is not None else None


def _add(a, b):
    return a + b if a is not None and b is not None else None


def _sub(a, *bs):
    if a is None:
        return None
    return a - sum(b for b in bs if b is not None)


def _days(a, b):
    v = _div(a, b)
    return v * 365 if v is not None else None


def compute(doc):
    figs = extract_figures(doc)
    ratios = []
    for group, name, unit, formula, fn, better in RATIOS:
        try:
            cy, py = fn(figs, 0), fn(figs, 1)
        except (TypeError, ZeroDivisionError):
            cy = py = None
        if cy is None and py is None:
            continue
        ratios.append({"group": group, "name": name, "unit": unit, "formula": formula, "cy": cy, "py": py,
                       "comment": _comment(name, unit, cy, py, better)})
    growth = []
    for key, label in (("revenue", "Revenue"), ("gross_profit", "Gross profit"), ("operating_profit", "Operating profit"),
                       ("profit", "Profit after tax"), ("total_assets", "Total assets"), ("total_equity", "Total equity"),
                       ("cfo", "Operating cash flow")):
        cy, py = _f(figs, key, 0), _f(figs, key, 1)
        if cy is not None and py not in (None, 0):
            growth.append({"name": label, "cy": cy, "py": py, "change": cy - py, "pct": (cy - py) / abs(py) * 100})
    return {"figures": {k: [float(v) if v is not None else None for v in pair] for k, pair in figs.items()},
            "ratios": ratios, "growth": growth,
            "horizontal": horizontal(doc), "highlights": _highlights(ratios, growth)}


def _fmt(v, unit):
    if v is None:
        return "n/a"
    if unit == "%":
        return f"{v:.1f}%"
    if unit == "days":
        return f"{v:.0f} days"
    if unit == "amount":
        return f"{v:,.0f}" if v >= 0 else f"({abs(v):,.0f})"
    return f"{v:.2f}x"


def _comment(name, unit, cy, py, better):
    if cy is None or py is None:
        return f"{name}: {_fmt(cy if cy is not None else py, unit)} (no comparison available)."
    diff = cy - py
    if abs(diff) < (0.05 if unit in ("times",) else 0.5) or (py and abs(diff) / abs(py) < 0.05):
        return f"{name} broadly unchanged at {_fmt(cy, unit)}."
    word = "increased" if diff > 0 else "decreased"
    tone = ""
    if better:
        good = (diff > 0) == (better == "up")
        tone = " — an improvement" if good else " — a deterioration; obtain an explanation"
    return f"{name} {word} from {_fmt(py, unit)} to {_fmt(cy, unit)}{tone}."


def _highlights(ratios, growth):
    out = []
    for g in growth:
        if abs(g["pct"]) >= 10:
            out.append(f"{g['name']} {'rose' if g['pct'] > 0 else 'fell'} {abs(g['pct']):.1f}% "
                       f"({g['py']:,.0f} → {g['cy']:,.0f}).")
    for r in ratios:
        if "deterioration" in r["comment"]:
            out.append(r["comment"])
    cur = next((r for r in ratios if r["name"] == "Current ratio"), None)
    if cur and cur["cy"] is not None and cur["cy"] < 1:
        out.append(f"Current ratio below 1 ({cur['cy']:.2f}x): consider the going concern assessment.")
    return out


def horizontal(doc):
    """Line-by-line change and common-size % for the statement of financial position and profit or loss."""
    out = []
    for t in doc.tables():
        if t.section not in ("SOFP", "SOCI") or not is_financial_table(t) or len(t.amount_cols) < 2:
            continue
        base = None
        for r in t.rows:
            if t.section == "SOFP" and r.label.lower() == "total assets":
                base = row_values(t, r)
            if t.section == "SOCI" and base is None and r.label and r.kind == "item":
                base = row_values(t, r)        # first line: revenue
        lines = []
        for r in t.rows:
            if r.kind not in ("item", "total", "grand_total", "subtotal"):
                continue
            v = row_values(t, r)
            cy, py = v[0], v[1]
            if cy is None and py is None:
                continue
            ch = (cy or 0) - (py or 0)
            pct = float(ch / abs(py) * 100) if py not in (None, Decimal(0)) else None
            common = float(cy / base[0] * 100) if base and base[0] not in (None, Decimal(0)) and cy is not None else None
            lines.append({"label": r.label or "(total)", "kind": r.kind, "cy": float(cy) if cy is not None else None,
                          "py": float(py) if py is not None else None, "change": float(ch), "pct": pct,
                          "common": common})
        if lines:
            out.append({"section": SECTIONS[t.section], "lines": lines})
    return out


# ----------------------------------------------------------------- charts

def _setup():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 9, "axes.edgecolor": GRID, "axes.labelcolor": MUTED,
        "xtick.color": MUTED, "ytick.color": MUTED, "axes.spines.top": False, "axes.spines.right": False,
        "axes.spines.left": False, "axes.grid": True, "axes.grid.axis": "y", "grid.color": GRID,
        "grid.linewidth": 0.6, "axes.titleweight": "bold", "axes.titlesize": 10.5, "axes.titlecolor": INK,
        "axes.titlelocation": "left", "legend.frameon": False,
    })
    return plt


def _bars(plt, ax, labels, cy, py, years, fmt=lambda v: f"{v:,.0f}"):
    import numpy as np
    x = np.arange(len(labels))
    w = 0.36
    for off, vals, color, name in ((-w / 2 - 0.01, py, PY_COLOR, years[1]), (w / 2 + 0.01, cy, CY_COLOR, years[0])):
        vals = [v if v is not None else 0 for v in vals]
        bars = ax.bar(x + off, vals, w, color=color, label=name, zorder=3)
        for b, v in zip(bars, vals):
            if v:
                ax.annotate(fmt(v), (b.get_x() + b.get_width() / 2, v), ha="center",
                            va="bottom" if v >= 0 else "top", fontsize=7.5, color=MUTED,
                            xytext=(0, 2 if v >= 0 else -2), textcoords="offset points")
    ax.set_xticks(x, labels)
    ax.axhline(0, color=MUTED, linewidth=0.8, zorder=2)
    ax.legend(loc="lower right", bbox_to_anchor=(1.0, 1.0), ncols=2, fontsize=8, borderaxespad=0.3)
    ax.margins(y=0.12)
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v:,.0f}"))


def charts(result, out_dir, years=("Current year", "Prior year")):
    """Write PNG charts; return [(title, path)]."""
    try:
        plt = _setup()
    except ImportError:
        return []
    f = result["figures"]
    get = lambda k, i: (f.get(k) or [None, None])[i]
    made = []

    def save(fig, name, title):
        path = os.path.join(out_dir, name)
        fig.tight_layout()
        fig.savefig(path, dpi=160)
        plt.close(fig)
        made.append((title, path))

    keys = [("revenue", "Revenue"), ("gross_profit", "Gross profit"), ("operating_profit", "Operating profit"),
            ("profit", "Profit after tax")]
    keys = [(k, l) for k, l in keys if get(k, 0) is not None or get(k, 1) is not None]
    if keys:
        fig, ax = plt.subplots(figsize=(7.2, 3.4))
        _bars(plt, ax, [l for _, l in keys], [get(k, 0) for k, _ in keys], [get(k, 1) for k, _ in keys], years)
        ax.set_title("Results of operations")
        save(fig, "chart-results.png", "Results of operations")

    margins = [r for r in result["ratios"] if r["group"] == "Profitability" and r["unit"] == "%"
               and r["name"] != "Effective tax rate"]
    if margins:
        fig, ax = plt.subplots(figsize=(7.2, 3.2))
        _bars(plt, ax, [{"Return on equity": "ROE", "Return on assets": "ROA"}.get(r["name"], r["name"].replace(" margin", ""))
               for r in margins],
              [r["cy"] for r in margins], [r["py"] for r in margins], years, fmt=lambda v: f"{v:.1f}%")
        ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v:.0f}%"))
        ax.set_title("Margins and returns")
        save(fig, "chart-margins.png", "Margins and returns")

    comp = [("total_non_current_assets", "Non-current assets"), ("total_current_assets", "Current assets"),
            ("total_equity", "Equity"), ("total_non_current_liabilities", "Non-current liabilities"),
            ("total_current_liabilities", "Current liabilities")]
    comp = [(k, l) for k, l in comp if get(k, 0) is not None or get(k, 1) is not None]
    if comp:
        fig, ax = plt.subplots(figsize=(7.2, 3.4))
        _bars(plt, ax, [l.replace(" ", "\n", 1) for _, l in comp], [get(k, 0) for k, _ in comp],
              [get(k, 1) for k, _ in comp], years)
        ax.set_title("Financial position: assets and how they are financed")
        save(fig, "chart-position.png", "Financial position")

    cf = [("cfo", "Operating"), ("cfi", "Investing"), ("cff", "Financing")]
    cf = [(k, l) for k, l in cf if get(k, 0) is not None or get(k, 1) is not None]
    if cf:
        fig, ax = plt.subplots(figsize=(7.2, 3.2))
        _bars(plt, ax, [l for _, l in cf], [get(k, 0) for k, _ in cf], [get(k, 1) for k, _ in cf], years,
              fmt=lambda v: f"{v:,.0f}" if v >= 0 else f"({abs(v):,.0f})")
        ax.set_title("Cash flows by activity")
        save(fig, "chart-cashflows.png", "Cash flows by activity")

    times = [r for r in result["ratios"] if r["unit"] == "times" and r["group"] in ("Liquidity", "Leverage")
             and r["name"] != "Interest cover"]          # cover runs on a much larger scale: table only
    if times:
        fig, ax = plt.subplots(figsize=(7.2, 3.2))
        _bars(plt, ax, [r["name"].replace(" ", "\n", 1) for r in times], [r["cy"] for r in times],
              [r["py"] for r in times], years, fmt=lambda v: f"{v:.2f}x")
        ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v:.1f}x"))
        ax.set_title("Liquidity and leverage")
        save(fig, "chart-liquidity.png", "Liquidity and leverage")

    days = [r for r in result["ratios"] if r["unit"] == "days"]
    if days:
        fig, ax = plt.subplots(figsize=(7.2, 3.0))
        _bars(plt, ax, [r["name"] for r in days], [r["cy"] for r in days], [r["py"] for r in days], years,
              fmt=lambda v: f"{v:.0f}")
        ax.set_title("Working capital cycle (days)")
        save(fig, "chart-days.png", "Working capital cycle")
    return made


# ----------------------------------------------------------------- report

def write_report(path, result, chart_files, settings, years):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    from .report import _cell, _shade, _widths

    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(2)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading("Financial Analysis", 0)
    d.add_paragraph(f"{settings.get('entity_name') or ''}  ·  {years[0]} compared with {years[1]}".strip(" ·"))
    if result["highlights"]:
        d.add_heading("Points for the reviewer", 1)
        for h in result["highlights"]:
            d.add_paragraph(h, style="List Bullet")
    for title, p in chart_files:
        d.add_picture(p, width=Cm(16.5))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_heading("Ratio analysis", 1)
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for c, h in enumerate(("Ratio", str(years[0]), str(years[1]), "Formula", "Observation")):
        _cell(t.rows[0].cells[c], h, bold=True)
        _shade(t.rows[0].cells[c], "D9E2F3")
    group = None
    for r in result["ratios"]:
        if r["group"] != group:
            group = r["group"]
            row = t.add_row().cells
            _cell(row[0], group, bold=True)
            for c in row:
                _shade(c, "F2F2F2")
        row = t.add_row().cells
        _cell(row[0], r["name"])
        _cell(row[1], _fmt(r["cy"], r["unit"]))
        _cell(row[2], _fmt(r["py"], r["unit"]))
        _cell(row[3], r["formula"], size=8)
        _cell(row[4], r["comment"], size=8)
    _widths(t, (Cm(3.6), Cm(2.0), Cm(2.0), Cm(4.4), Cm(5.5)))
    if result["growth"]:
        d.add_heading("Growth", 1)
        t = d.add_table(rows=1, cols=5)
        t.style = "Table Grid"
        for c, h in enumerate(("Item", str(years[0]), str(years[1]), "Change", "%")):
            _cell(t.rows[0].cells[c], h, bold=True)
            _shade(t.rows[0].cells[c], "D9E2F3")
        for g in result["growth"]:
            row = t.add_row().cells
            for c, v in enumerate((g["name"], _fmt(g["cy"], "amount"), _fmt(g["py"], "amount"),
                                   _fmt(g["change"], "amount"), f"{g['pct']:+.1f}%")):
                _cell(row[c], v)
        _widths(t, (Cm(4.5), Cm(3.2), Cm(3.2), Cm(3.2), Cm(2.4)))
    for block in result["horizontal"]:
        d.add_heading(f"Horizontal and vertical analysis — {block['section']}", 1)
        t = d.add_table(rows=1, cols=6)
        t.style = "Table Grid"
        base = "% of total assets" if "Position" in block["section"] else "% of revenue"
        for c, h in enumerate(("Line", str(years[0]), str(years[1]), "Change", "Change %", base)):
            _cell(t.rows[0].cells[c], h, bold=True)
            _shade(t.rows[0].cells[c], "D9E2F3")
        for ln in block["lines"]:
            row = t.add_row().cells
            bold = ln["kind"] in ("total", "grand_total")
            vals = (ln["label"], _fmt(ln["cy"], "amount"), _fmt(ln["py"], "amount"), _fmt(ln["change"], "amount"),
                    f"{ln['pct']:+.1f}%" if ln["pct"] is not None else "—",
                    f"{ln['common']:.1f}%" if ln["common"] is not None else "—")
            for c, v in enumerate(vals):
                _cell(row[c], v, bold=bold, size=8.5)
            if ln["pct"] is not None and abs(ln["pct"]) >= 20 and ln["kind"] == "item" and abs(ln["change"]) > 0:
                _shade(row[4], "FFF2CC")
        _widths(t, (Cm(5.4), Cm(2.4), Cm(2.4), Cm(2.4), Cm(1.9), Cm(2.4)))
        d.add_paragraph("Shaded: lines that moved by 20% or more — obtain explanations.").runs[0].italic = True
    d.save(path)


def add_to_workbook(xlsx_path, result, years):
    """A 'Ratios' sheet and an 'Analysis' sheet in the Excel output."""
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.load_workbook(xlsx_path, keep_vba=xlsx_path.lower().endswith(".xlsm"))
    for name in ("Ratios", "Analysis"):
        if name in wb.sheetnames:
            del wb[name]
    ws = wb.create_sheet("Ratios")
    ws.append(["Ratio", str(years[0]), str(years[1]), "Formula", "Observation"])
    for c in ws[1]:
        c.font = Font(bold=True)
    group = None
    for r in result["ratios"]:
        if r["group"] != group:
            group = r["group"]
            ws.append([group])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.append([r["name"], _round(r["cy"]), _round(r["py"]), r["formula"], r["comment"]])
        fmt = '0.0"%"' if r["unit"] == "%" else '0" days"' if r["unit"] == "days" else \
            '#,##0;(#,##0)' if r["unit"] == "amount" else '0.00"x"'
        for c in (2, 3):
            ws.cell(ws.max_row, c).number_format = fmt
    for col, w in zip("ABCDE", (28, 12, 12, 48, 70)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("Analysis")
    for block in result["horizontal"]:
        ws2.append([block["section"]])
        ws2.cell(ws2.max_row, 1).font = Font(bold=True)
        ws2.append(["Line", str(years[0]), str(years[1]), "Change", "Change %", "Common size %"])
        for c in ws2[ws2.max_row]:
            c.font = Font(bold=True)
        for ln in block["lines"]:
            ws2.append([ln["label"], ln["cy"], ln["py"], ln["change"],
                        _round(ln["pct"]), _round(ln["common"])])
            for c in (2, 3, 4):
                ws2.cell(ws2.max_row, c).number_format = "#,##0;(#,##0);-"
            for c in (5, 6):
                ws2.cell(ws2.max_row, c).number_format = '0.0"%"'
        ws2.append([])
    for col, w in zip("ABCDEF", (44, 14, 14, 14, 11, 14)):
        ws2.column_dimensions[col].width = w
    wb.save(xlsx_path)


def _round(v):
    return round(v, 2) if v is not None else None
