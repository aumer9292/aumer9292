"""Formatting templates and the review settings the user chooses."""

import copy

PRESETS = {
    "classic": {
        "label": "Classic (Times New Roman, audit-file style)",
        "font": "Times New Roman", "size": 11, "table_size": 10, "color": "000000",
        "heading_font": "Times New Roman", "title_size": 14, "h1_size": 12, "h2_size": 11,
        "heading_color": "000000", "heading_caps": "upper_titles",
    },
    "modern": {
        "label": "Modern (Arial, clean corporate)",
        "font": "Arial", "size": 10, "table_size": 9, "color": "1F1F1F",
        "heading_font": "Arial", "title_size": 14, "h1_size": 11, "h2_size": 10,
        "heading_color": "1F3864", "heading_caps": "upper_titles",
    },
    "calibri": {
        "label": "Calibri (Word default family)",
        "font": "Calibri", "size": 11, "table_size": 10, "color": "000000",
        "heading_font": "Calibri", "title_size": 14, "h1_size": 12, "h2_size": 11,
        "heading_color": "000000", "heading_caps": "upper_titles",
    },
    "garamond": {
        "label": "Annual report (Garamond)",
        "font": "Garamond", "size": 11, "table_size": 10, "color": "000000",
        "heading_font": "Garamond", "title_size": 15, "h1_size": 12, "h2_size": 11,
        "heading_color": "000000", "heading_caps": "upper_titles",
    },
}

DEFAULTS = {
    # engagement details
    "entity_name": "",
    "framework": "IFRS",
    "jurisdiction": "",
    "entity_type": "Private company",
    "financial_year": "",
    "currency": "",
    "unit": "",                    # e.g. "'000"
    # page
    "template": "classic",
    "page_size": "A4",             # A4 or Letter
    "margin_top_cm": 2.5, "margin_bottom_cm": 2.5, "margin_left_cm": 2.5, "margin_right_cm": 2.0,
    "line_spacing": 1.0,
    "space_after_pt": 6,
    "justify_body": True,
    # figures
    "negative_style": "brackets",  # brackets or minus
    "zero_style": "-",            # '-' or '0'
    "reformat_figures": True,
    "table_borders": "financial",  # financial (rules on totals only) or keep
    # structure
    "statement_page_breaks": True,
    "header_footer": True,
    "page_numbers": True,
    # proofreading
    "apply_spelling": False,
    "apply_spacing_fixes": True,
    "custom_words": [],
    "ai_review": False,
    # outputs
    "outputs": ["docx", "xlsx", "pdf", "report"],
}


def build_settings(user=None, reference_style=None):
    s = copy.deepcopy(DEFAULTS)
    if user:
        for k, v in user.items():
            if v is None or v == "":
                continue
            s[k] = v
    preset = PRESETS.get(s.get("template"), PRESETS["classic"])
    style = dict(preset)
    if reference_style:
        style.update({k: v for k, v in reference_style.items() if v})
        for k in ("margin_top_cm", "margin_bottom_cm", "margin_left_cm", "margin_right_cm", "page_size"):
            if reference_style.get(k):
                s[k] = reference_style[k]
    s["style"] = style
    for k in ("margin_top_cm", "margin_bottom_cm", "margin_left_cm", "margin_right_cm", "line_spacing",
              "space_after_pt"):
        s[k] = float(s[k])
    for k in ("apply_spelling", "apply_spacing_fixes", "reformat_figures", "statement_page_breaks",
              "header_footer", "page_numbers", "justify_body", "ai_review"):
        if isinstance(s[k], str):
            s[k] = s[k].lower() in ("1", "true", "yes", "on")
    if isinstance(s["custom_words"], str):
        s["custom_words"] = [w.strip() for w in s["custom_words"].replace(",", "\n").splitlines() if w.strip()]
    return s


def style_from_reference(path):
    """Read the house style from a previously approved Word document."""
    import docx
    from docx.shared import Length
    d = docx.Document(path)
    out = {}

    def font_of(style_name):
        try:
            st = d.styles[style_name]
        except KeyError:
            return None, None
        return st.font.name, (st.font.size.pt if st.font.size else None)

    name, size = font_of("Normal")
    # Direct formatting often overrides the Normal style: take the most common run font.
    counts, sizes = {}, {}
    for p in d.paragraphs[:400]:
        for r in p.runs:
            if r.text.strip():
                if r.font.name:
                    counts[r.font.name] = counts.get(r.font.name, 0) + len(r.text)
                if r.font.size:
                    sizes[r.font.size.pt] = sizes.get(r.font.size.pt, 0) + len(r.text)
    if counts:
        name = max(counts, key=counts.get)
    if sizes:
        size = max(sizes, key=sizes.get)
    if name:
        out["font"] = name
        out["heading_font"] = name
    if size:
        out["size"] = size
        out["table_size"] = max(8, size - 1)
    h1, h1s = font_of("Heading 1")
    if h1s:
        out["h1_size"] = h1s
    tsz = []
    for t in d.tables[:10]:
        for row in t.rows[:20]:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        if r.text.strip() and r.font.size:
                            tsz.append(r.font.size.pt)
    if tsz:
        out["table_size"] = max(set(tsz), key=tsz.count)
    sec = d.sections[0]
    cm = lambda v: round(Length(v).cm, 2) if v is not None else None
    out["margin_top_cm"] = cm(sec.top_margin)
    out["margin_bottom_cm"] = cm(sec.bottom_margin)
    out["margin_left_cm"] = cm(sec.left_margin)
    out["margin_right_cm"] = cm(sec.right_margin)
    if sec.page_width is not None:
        out["page_size"] = "Letter" if abs(Length(sec.page_width).inches - 8.5) < 0.1 else "A4"
    return out
