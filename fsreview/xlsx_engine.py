"""Excel output.

Word, PDF and text inputs get a workbook with one worksheet per financial table,
figures stored as numbers. (An Excel input is never re-saved: see xlsx_patch.)
"""

import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .classify import GRAND_TOTAL, PRIMARY, SECTIONS, TOTAL, is_financial_table
from .numbers import YEAR, parse_amount

THIN = Side(style="thin", color="000000")
DOUBLE = Side(style="double", color="000000")


def _num_format(decimals, settings):
    d = "." + "0" * decimals if decimals else ""
    zero = '"-"' if settings.get("zero_style", "-") == "-" else "0" + d
    if settings.get("negative_style", "brackets") == "brackets":
        return f"#,##0{d}_);(#,##0{d});{zero}_)"
    return f"#,##0{d};-#,##0{d};{zero}"


def _is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _print_setup(ws, settings, changes):
    """Print the sheet the way a signed set is printed: A4/Letter, one page wide, centred, page numbers."""
    ps = ws.page_setup
    ps.paperSize = ws.PAPERSIZE_LETTER if settings.get("page_size") == "Letter" else ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ps.fitToWidth = 1
    ps.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    m = ws.page_margins
    m.left, m.right, m.top, m.bottom = 0.7, 0.5, 0.75, 0.75
    m.left += (settings.get("gutter_cm") or 0) / 2.54
    if not (ws.oddFooter.center.text or ws.oddFooter.left.text or ws.oddFooter.right.text):
        ws.oddFooter.center.text = "&P"
        ws.oddFooter.center.size = 8
    ws.sheet_view.showGridLines = False
    changes.add("Excel print settings set (paper, fit to page width, centred, page numbers)")


def _format_sheet(ws, vws, settings, st, changes, bounds=None):
    if ws.max_row < 1 or ws.max_column < 1:
        return
    c0, r0, c1, r1 = bounds or (1, 1, min(ws.max_column, 60), min(ws.max_row, 5000))
    rows = [r for r in range(r0, r1 + 1) if not ws.row_dimensions[r].hidden]
    columns = [c for c in range(c0, c1 + 1) if not ws.column_dimensions[get_column_letter(c)].hidden]
    # label column: the column holding the most text
    label_col, best = columns[0] if columns else 1, -1
    for c in columns[:6]:
        n = sum(1 for r in rows if isinstance(vws.cell(r, c).value, str) and len(vws.cell(r, c).value.strip()) > 3)
        if n > best:
            label_col, best = c, n
    amount_cols = []
    for c in columns:
        if c == label_col:
            continue
        cells = [ws.cell(r, c) for r in rows if ws.cell(r, c).value is not None]
        nums = [x for x in cells if _is_number(x.value) or (isinstance(x.value, str) and x.value.startswith("="))]
        nums = [x for x in nums if not (isinstance(x.value, int) and 1900 <= x.value <= 2100)]
        if nums and len(nums) >= 0.5 * len(cells) and not all(_is_number(vws.cell(x.row, x.column).value)
                                                               and abs(vws.cell(x.row, x.column).value or 0) < 100
                                                               for x in nums):
            amount_cols.append(c)
    first_data = (rows[-1] + 1) if rows else r1 + 1
    for r in rows:
        if any(_is_number(vws.cell(r, c).value) and not (isinstance(vws.cell(r, c).value, int)
               and 1900 <= vws.cell(r, c).value <= 2100) for c in amount_cols) or \
                any(isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.startswith("=") for c in amount_cols):
            first_data = r
            break
    decimals = {}
    for c in amount_cols:
        # keep the decimals the sheet already shows; never reveal hidden decimals
        decs = [len(m.group(1)) for r in rows if r >= first_data
                for m in [re.search(r"#,##0\.(0+)", (ws.cell(r, c).number_format or "").split(";")[0])] if m]
        decimals[c] = max(set(decs), key=decs.count) if decs else 0

    font = st["font"]
    size = st["table_size"]
    prev_kind = None
    for r in rows:
        label = str(vws.cell(r, label_col).value or "").strip()
        has_amt = any(vws.cell(r, c).value is not None or ws.cell(r, c).value is not None for c in amount_cols)
        formulas = [str(ws.cell(r, c).value) for c in amount_cols if isinstance(ws.cell(r, c).value, str)
                    and ws.cell(r, c).value.startswith("=")]
        if r < first_data:
            kind = "header" if has_amt or r < first_data else "heading"
        elif not label and not has_amt:
            kind = "blank"
        elif not has_amt:
            kind = "heading"
        elif GRAND_TOTAL.search(label):
            kind = "grand_total"
        elif TOTAL.search(label) or any(re.search(r"SUM\(", f, re.I) for f in formulas) or (not label and has_amt):
            kind = "total"
        else:
            kind = "item"
        for c in columns:
            cell = ws.cell(r, c)
            if cell.value is None and c not in amount_cols:
                continue
            f = cell.font
            bold = f.bold
            if kind == "header" or (kind in ("grand_total", "heading") and (c == label_col or c in amount_cols)):
                bold = True
            elif kind == "item" and c in amount_cols:
                bold = False
            cell.font = Font(name=font, size=size, bold=bold, italic=f.italic, underline=f.underline,
                             color=f.color if f.color is not None and f.color.type == "rgb" and f.color.rgb not in
                             (None, "FF000000") else st["color"])
            if c in amount_cols:
                fmt = cell.number_format or "General"
                is_pct_or_date = "%" in fmt or re.search(r"[dmy]{2,}", fmt, re.I)
                if kind == "header":
                    cell.alignment = Alignment(horizontal="right", vertical="bottom", wrap_text=True)
                elif not is_pct_or_date and (_is_number(cell.value) or (isinstance(cell.value, str)
                                                                         and cell.value.startswith("="))):
                    new = _num_format(decimals.get(c, 0), settings)
                    has_style = "#,##0" in fmt and ("(" in fmt or "-" in fmt)
                    if bounds and has_style:
                        pass        # the sheet already presents figures with separators and negatives: keep it
                    elif cell.number_format != new:
                        cell.number_format = new
                        changes.add("Excel number formats standardised")
                    cell.alignment = Alignment(horizontal="right", vertical="bottom")
                if kind in ("total", "grand_total") and prev_kind in ("item", "total") and cell.value is not None:
                    cell.border = Border(top=THIN, bottom=DOUBLE if kind == "grand_total" else None)
                elif kind == "grand_total" and cell.value is not None:
                    cell.border = Border(bottom=DOUBLE)
        if kind != "blank":
            prev_kind = kind
    if bounds is None:
        # a sheet this tool built: size the columns. An existing sheet keeps its tuned print layout.
        ws.column_dimensions[get_column_letter(label_col)].width = min(60, max(
            [len(str(vws.cell(r, label_col).value or "")) for r in rows] + [20]) + 2)
        for c in amount_cols:
            ws.column_dimensions[get_column_letter(c)].width = 15 if decimals.get(c) else 14
        _print_setup(ws, settings, changes)
    changes.add("Excel worksheets formatted")


def model_to_workbook(model, out, settings, changes):
    """Workbook with one worksheet per financial table (from Word, PDF or text input)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    st = settings["style"]
    used = set()
    count = 0
    last_title = None
    for b in model.blocks:
        if b.kind == "paragraph" and b.role == "title":
            last_title = b.text.strip()
        if b.kind != "table" or not is_financial_table(b):
            continue
        count += 1
        base = {"SOFP": "SOFP", "SOCI": "SOCI", "SOCE": "SOCE", "SOCF": "SOCF"}.get(b.section)
        if not base:
            base = f"Note {b.note_number}" if b.note_number else f"Table {count}"
        name = base[:28]
        k = 2
        while name in used:
            name = f"{base[:25]} ({k})"
            k += 1
        used.add(name)
        ws = wb.create_sheet(name)
        r0 = 1
        heading = SECTIONS.get(b.section, "") if b.section in PRIMARY else (last_title or "")
        if heading:
            ws.cell(1, 1, heading).font = Font(name=st["font"], size=st["table_size"] + 2, bold=True)
            r0 = 3
        for ri, row in enumerate(b.rows):
            for ci, txt in enumerate(row.cells):
                v = parse_amount(txt) if ci in b.amount_cols and not YEAR.match(txt.strip() or "x") else None
                if row.kind == "header":
                    v = None
                cell = ws.cell(r0 + ri, ci + 1)
                if v is not None and txt.strip():
                    cell.value = float(v) if v != v.to_integral_value() else int(v)
                elif txt:
                    cell.value = txt
        _format_sheet(ws, ws, settings, st, changes)
    if not wb.worksheets:
        ws = wb.create_sheet("No tables")
        ws["A1"] = "No financial tables were found in the source document."
    wb.save(out)
