"""The review report: a Word document for the file and JSON for the screen."""

import datetime as _dt
import json

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .classify import SECTIONS

SEV_ORDER = {"error": 0, "warning": 1, "info": 2}
CATEGORY_TITLES = [
    ("integrity", "Content integrity"),
    ("disclosure", "Required disclosures"),
    ("accounting", "Accounting cross-checks"),
    ("casting", "Casting (arithmetic)"),
    ("consistency", "Consistency and comparatives"),
    ("proofreading", "Proofreading"),
    ("extraction", "Extraction and items for manual review"),
    ("formatting", "Formatting"),
]
SEV_COLOR = {"error": "C00000", "warning": "B26B00", "info": "2F5496"}


def summary(issues):
    out = {"error": 0, "warning": 0, "info": 0, "passed": 0}
    for i in issues:
        if i.status == "passed":
            out["passed"] += 1
        else:
            out[i.severity] += 1
    return out


def sections_found(model):
    seen = []
    for b in model.blocks:
        if b.section and b.section not in seen and b.section not in ("COVER",):
            seen.append(b.section)
    return [SECTIONS.get(s, s) for s in seen]


def to_json(path, model, settings, issues, changes, outputs, notes, extra=None):
    data = {
        "source": model.source_name,
        "generated": _dt.datetime.now().isoformat(timespec="seconds"),
        "settings": {k: v for k, v in settings.items() if k not in ("style",)},
        "summary": summary(issues),
        "sections": sections_found(model),
        "issues": [i.as_dict() for i in sorted(issues, key=_key)],
        "changes": {"counts": changes.counts,
                    "edits": [{"location": l, "before": b, "after": a, "reason": r}
                              for l, b, a, r in changes.text[:2000]]},
        "outputs": outputs,
        "notes": notes,
        **(extra or {}),
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    return data


def _key(i):
    order = [c for c, _ in CATEGORY_TITLES]
    return (order.index(i.category) if i.category in order else 99, i.status == "passed", SEV_ORDER[i.severity])


def _shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _widths(table, widths):
    """Fixed column widths that Word and LibreOffice both respect."""
    table.autofit = False
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w / 635)))
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = w


def _cell(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(0)


def to_docx(path, model, settings, issues, changes, outputs, notes, extra=None):
    extra = extra or {}
    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(2)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    d.add_heading("Financial Statements Review Report", 0)
    meta = [("Source document", model.source_name),
            ("Entity", settings.get("entity_name") or "—"),
            ("Framework", settings.get("framework") or "—"),
            ("Jurisdiction", settings.get("jurisdiction") or "—"),
            ("Entity category", f"{extra['category']['name']} ({extra['category']['code']})" if extra.get("category") else "—"),
            ("Financial year", str(settings.get("financial_year") or "—")),
            ("Presentation currency", " ".join(x for x in (settings.get("currency"), settings.get("unit")) if x) or "—"),
            ("Template", settings["style"].get("label", settings.get("template"))),
            ("Prepared", _dt.datetime.now().strftime("%d %B %Y %H:%M"))]
    t = d.add_table(rows=len(meta), cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate(meta):
        _cell(t.cell(i, 0), k, bold=True)
        _shade(t.cell(i, 0), "F2F2F2")
        _cell(t.cell(i, 1), v)
    _widths(t, (Cm(5), Cm(12.5)))

    s = summary(issues)
    d.add_heading("Summary", 1)
    p = d.add_paragraph()
    for label, key in (("Errors", "error"), ("Warnings", "warning"), ("For information", "info"),
                       ("Checks passed", "passed")):
        r = p.add_run(f"{label}: {s[key]}    ")
        r.bold = True
        if key in SEV_COLOR:
            r.font.color.rgb = RGBColor.from_string(SEV_COLOR[key])
    if s["error"]:
        d.add_paragraph("The financial statements should not be finalised until the errors below are resolved "
                        "or explained.")
    found = sections_found(model)
    d.add_paragraph("Sections identified: " + (", ".join(found) if found else "none"))
    for n in notes:
        d.add_paragraph(n, style="List Bullet")

    for cat, title in CATEGORY_TITLES:
        rows = sorted([i for i in issues if i.category == cat], key=_key)
        if not rows:
            continue
        d.add_heading(title, 1)
        tbl = d.add_table(rows=1, cols=4)
        tbl.style = "Table Grid"
        widths = (Cm(1.9), Cm(8.6), Cm(3.4), Cm(3.6))
        for c, h in enumerate(("Status", "Finding", "Where", "Suggested action")):
            _cell(tbl.rows[0].cells[c], h, bold=True)
            _shade(tbl.rows[0].cells[c], "D9E2F3")
        for i in rows:
            cells = tbl.add_row().cells
            status = "Passed" if i.status == "passed" else i.severity.title()
            color = "2E7D32" if i.status == "passed" else SEV_COLOR[i.severity]
            _cell(cells[0], status, bold=True, color=color)
            _cell(cells[1], i.message)
            _cell(cells[2], ", ".join(x for x in (i.section, i.location) if x))
            _cell(cells[3], i.suggestion)
        _widths(tbl, widths)

    if extra.get("checklist"):
        d.add_heading(f"Disclosure checklist — {extra['category']['name']}", 1)
        d.add_paragraph(f"Framework: {extra['category']['framework']}. Disclosure requirements: "
                        f"{extra['category']['schedule']}. The checklist confirms that each disclosure is present; "
                        "its content still needs review.")
        tbl = d.add_table(rows=1, cols=4)
        tbl.style = "Table Grid"
        for c, h in enumerate(("Disclosure", "Reference", "Status", "Found")):
            _cell(tbl.rows[0].cells[c], h, bold=True)
            _shade(tbl.rows[0].cells[c], "D9E2F3")
        label = {"present": ("Present", "2E7D32"), "missing": ("Not found", "C00000"), "n/a": ("Not applicable", "7F7F7F")}
        for r in extra["checklist"]:
            cells = tbl.add_row().cells
            _cell(cells[0], r["title"])
            _cell(cells[1], r["ref"], size=8)
            _cell(cells[2], label[r["status"]][0], bold=True, color=label[r["status"]][1])
            _cell(cells[3], r["evidence"], size=8)
        _widths(tbl, (Cm(6.2), Cm(3.6), Cm(2.2), Cm(5.5)))

    d.add_heading("Formatting applied", 1)
    if changes.counts:
        for what, n in sorted(changes.counts.items()):
            d.add_paragraph(f"{what}: {n}", style="List Bullet")
    else:
        d.add_paragraph("No formatting changes were needed.")
    if changes.text:
        d.add_heading("Text changes (audit trail)", 2)
        d.add_paragraph("Every change to wording or to the way a figure is written is listed below. Figure "
                        "changes alter presentation only; the value is verified to be identical.")
        tbl = d.add_table(rows=1, cols=4)
        tbl.style = "Table Grid"
        for c, h in enumerate(("Where", "Before", "After", "Reason")):
            _cell(tbl.rows[0].cells[c], h, bold=True)
            _shade(tbl.rows[0].cells[c], "D9E2F3")
        for loc, before, after, reason in changes.text[:1500]:
            cells = tbl.add_row().cells
            _cell(cells[0], loc, size=8)
            _cell(cells[1], before[:300], size=8)
            _cell(cells[2], after[:300], size=8)
            _cell(cells[3], reason, size=8)
        _widths(tbl, (Cm(3.6), Cm(5.5), Cm(5.5), Cm(2.9)))
        if len(changes.text) > 1500:
            d.add_paragraph(f"… and {len(changes.text) - 1500} more (see the JSON report).")
    d.add_heading("Output files", 1)
    for k, v in outputs.items():
        d.add_paragraph(f"{k}: {v}", style="List Bullet")
    d.add_paragraph().add_run("This report was produced automatically. It supports, and does not replace, the "
                              "professional judgement of the preparer and reviewer.").italic = True
    d.save(path)
